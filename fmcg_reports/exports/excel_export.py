"""
Excel Export — OpenPyXL. Company name pulled from DB dynamically.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import send_file
import io
from datetime import datetime

HEADER_BG  = 'FF1A4E8F'
HEADER_FT  = 'FFFFFFFF'
ACCENT_BG  = 'FFF4A621'
TOTAL_BG   = 'FFDCE8F8'
ALT_ROW_BG = 'FFF7F9FC'


def thin_border():
    s = Side(style='thin', color='FFD0D0D0')
    return Border(left=s, right=s, top=s, bottom=s)


def money(val):
    try:    return float(val)
    except: return val


def _get_company_name():
    try:
        import pyodbc
        from app import DB_SERVER, DB_NAME, DB_USER, DB_PASSWORD
        conn = pyodbc.connect(
            f"DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={DB_SERVER};"
            f"DATABASE={DB_NAME};UID={DB_USER};PWD={DB_PASSWORD};TrustServerCertificate=yes;",
            timeout=10)
        cursor = conn.cursor()
        cursor.execute("SELECT company_name FROM dbo.company_master")
        row = cursor.fetchone()
        conn.close()
        return row.company_name if row else 'Business Reports'
    except Exception:
        return 'Business Reports'


def generate_excel(df, title, report_key):
    company_name = _get_company_name()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Report'

    # Title row
    n_cols = len(df.columns)
    ws.merge_cells(f'A1:{get_column_letter(n_cols)}1')
    tc = ws['A1']
    tc.value     = f'{company_name} — {title}'
    tc.font      = Font(name='Calibri', bold=True, size=14, color=HEADER_FT)
    tc.fill      = PatternFill('solid', fgColor=HEADER_BG)
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f'A2:{get_column_letter(n_cols)}2')
    sc = ws['A2']
    sc.value     = f'Generated: {datetime.now().strftime("%d/%m/%Y %I:%M %p")}   |   Records: {len(df)}'
    sc.font      = Font(name='Calibri', italic=True, size=9, color='FF555555')
    sc.alignment = Alignment(horizontal='center')

    # Header row
    header_row = 3
    for ci, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=ci, value=col_name.replace('_', ' '))
        cell.font      = Font(name='Calibri', bold=True, size=10, color=HEADER_FT)
        cell.fill      = PatternFill('solid', fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = thin_border()
    ws.row_dimensions[header_row].height = 25

    # Data rows
    numeric_cols = df.select_dtypes(include='number').columns.tolist()
    date_cols    = [c for c in df.columns if 'DATE' in c.upper() or 'date' in c.lower()]

    for ri, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        is_alt = (ri - header_row) % 2 == 0
        for ci, (col_name, val) in enumerate(row.items(), start=1):
            cell = ws.cell(row=ri, column=ci)
            if col_name in numeric_cols:
                try:    cell.value = float(val)
                except: cell.value = val
                cell.number_format = '#,##0.00'
                cell.alignment     = Alignment(horizontal='right')
            elif col_name in date_cols:
                cell.value     = str(val) if val is not None else ''
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.value     = str(val) if val is not None else ''
                cell.alignment = Alignment(horizontal='left')
            cell.font   = Font(name='Calibri', size=9)
            cell.fill   = PatternFill('solid', fgColor=ALT_ROW_BG if is_alt else 'FFFFFFFF')
            cell.border = thin_border()

    # Totals row
    if not df.empty and numeric_cols:
        tr = ws.max_row + 1
        ws.cell(row=tr, column=1, value='TOTAL').font = Font(bold=True, name='Calibri', size=10)
        ws.cell(row=tr, column=1).fill = PatternFill('solid', fgColor=TOTAL_BG)
        for ci, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=tr, column=ci)
            if col_name in numeric_cols:
                cl = get_column_letter(ci)
                cell.value         = f'=SUM({cl}{header_row+1}:{cl}{tr-1})'
                cell.number_format = '#,##0.00'
                cell.font          = Font(bold=True, name='Calibri', size=10)
                cell.alignment     = Alignment(horizontal='right')
            cell.fill   = PatternFill('solid', fgColor=TOTAL_BG)
            cell.border = thin_border()

    # Auto column width
    for ci, col_name in enumerate(df.columns, start=1):
        cl      = get_column_letter(ci)
        max_len = max(len(str(col_name)),
                      df[col_name].astype(str).str.len().max() if not df.empty else 0)
        ws.column_dimensions[cl].width = min(max(max_len + 2, 10), 35)

    ws.freeze_panes    = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    if not df.empty and numeric_cols:
        ws2 = wb.create_sheet('Summary')
        ws2['A1']      = f'{company_name} — Summary Statistics'
        ws2['A1'].font = Font(bold=True, size=12, color=HEADER_FT)
        ws2['A1'].fill = PatternFill('solid', fgColor=HEADER_BG)
        ws2.merge_cells('A1:D1')
        for hi, h in enumerate(['Column', 'Total', 'Average', 'Max'], start=1):
            c      = ws2.cell(row=2, column=hi, value=h)
            c.font = Font(bold=True, color=HEADER_FT)
            c.fill = PatternFill('solid', fgColor=ACCENT_BG)
        for ri2, col_name in enumerate(numeric_cols, start=3):
            ws2.cell(row=ri2, column=1, value=col_name.replace('_', ' '))
            ws2.cell(row=ri2, column=2, value=round(float(df[col_name].sum()), 2)).number_format = '#,##0.00'
            ws2.cell(row=ri2, column=3, value=round(float(df[col_name].mean()), 2)).number_format = '#,##0.00'
            ws2.cell(row=ri2, column=4, value=round(float(df[col_name].max()), 2)).number_format = '#,##0.00'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{report_key}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)
